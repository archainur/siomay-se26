"""Integration-level checks that generators pass complete source URLs onward."""

import io
import unittest
from unittest.mock import patch

import pandas as pd
from docx import Document
from PIL import Image

from src import bast, bapp_pml, bapp_pml_t2, bapp_ppl, bapp_ppl_t2, bukti_terima
from src.document_generator import insert_custom_url_images


REMOTE_URL = (
    "https://progres-mitra.bpskaro.site/"
    "api/dokumen-mitra/62/ss-fasih"
)


def _image_result(size=(120, 80)):
    image = Image.new("RGB", size, color="navy")
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    return stream, image


def _evidence_result(_source):
    stream, image = _image_result()
    size = image.size
    image.close()
    return [("image", stream, size)]


def _document_with_placeholder(placeholder):
    document = Document()
    paragraph = document.add_paragraph(placeholder)
    return document, paragraph


class BuiltinRemoteEvidenceTests(unittest.TestCase):
    def test_bapp_termin1_ppl_and_pml_accept_extensionless_source_url(self):
        for module in (bapp_ppl, bapp_pml):
            with self.subTest(module=module.__name__):
                document, paragraph = _document_with_placeholder(
                    module.BUKTI_PLACEHOLDER
                )
                with patch.object(
                    module,
                    "_download_drive_image",
                    side_effect=lambda _source: _image_result(),
                ) as downloader:
                    count, warnings = module.insert_evidence_images(
                        document, REMOTE_URL
                    )

                self.assertEqual(count, 1)
                self.assertEqual(warnings, [])
                downloader.assert_called_once_with(REMOTE_URL)
                self.assertEqual(len(document.inline_shapes), 1)
                self.assertEqual(paragraph.text, "")

    def test_bapp_termin2_ppl_and_pml_accept_extensionless_source_url(self):
        for module in (bapp_ppl_t2, bapp_pml_t2):
            with self.subTest(module=module.__name__):
                document, paragraph = _document_with_placeholder(
                    module.BUKTI_PLACEHOLDER
                )
                with patch.object(
                    module, "_download_drive_evidence", side_effect=_evidence_result
                ) as downloader:
                    count, warnings = module.insert_evidence_images(
                        document, REMOTE_URL
                    )

                self.assertEqual(count, 1)
                self.assertEqual(warnings, [])
                downloader.assert_called_once_with(REMOTE_URL)
                self.assertEqual(len(document.inline_shapes), 1)
                self.assertEqual(paragraph.text, "")

    def test_bast_reuses_the_termin2_universal_evidence_path(self):
        document, paragraph = _document_with_placeholder(bast.BUKTI_PLACEHOLDER)
        with patch.object(
            bapp_ppl_t2, "_download_drive_evidence", side_effect=_evidence_result
        ) as downloader:
            count, warnings = bast.insert_evidence_images(
                document, REMOTE_URL, placeholder=bast.BUKTI_PLACEHOLDER
            )

        self.assertEqual(count, 1)
        self.assertEqual(warnings, [])
        downloader.assert_called_once_with(REMOTE_URL)
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertEqual(paragraph.text, "")

    def test_bukti_terima_accepts_extensionless_source_url(self):
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        with patch.object(
            bukti_terima,
            "_download_drive_image",
            side_effect=lambda _source: _image_result(),
        ) as downloader:
            warnings = bukti_terima._fill_person_cell(
                cell,
                {
                    "nik": "6304000000000001",
                    "nama_lengkap": "PETUGAS UJI",
                    "wilayah_tugas": "Wilayah",
                    "operator": "Operator",
                    "no_telp": "0800000000",
                    "link_bukti_terima": REMOTE_URL,
                },
                0,
            )

        self.assertEqual(warnings, [])
        downloader.assert_called_once_with(REMOTE_URL)
        self.assertEqual(len(cell._tc.xpath(".//w:drawing")), 1)

    def test_bukti_terima_pdf_warning_does_not_abort_cell_batch(self):
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        with patch.object(
            bukti_terima,
            "_download_drive_image",
            side_effect=RuntimeError(
                "Resource berhasil diunduh tetapi bukan gambar yang didukung."
            ),
        ):
            warnings = bukti_terima._fill_person_cell(
                cell,
                {
                    "nik": "6304000000000001",
                    "nama_lengkap": "PETUGAS UJI",
                    "wilayah_tugas": "Wilayah",
                    "operator": "Operator",
                    "no_telp": "0800000000",
                    "link_bukti_terima": REMOTE_URL + ".pdf",
                },
                0,
            )

        self.assertEqual(len(warnings), 1)
        self.assertIn("Gagal memuat gambar", warnings[0])
        self.assertEqual(len(cell._tc.xpath(".//w:drawing")), 0)

    def test_custom_placeholder_uses_the_same_extensionless_source_contract(self):
        document, paragraph = _document_with_placeholder("Bukti: {{ss_fasih}}")
        row = pd.Series({"ss_fasih": REMOTE_URL})

        with patch(
            "utils.images.download_url_evidence", side_effect=_evidence_result
        ) as downloader:
            consumed = insert_custom_url_images(document, row)

        self.assertEqual(consumed, {"{{ss_fasih}}"})
        downloader.assert_called_once_with(REMOTE_URL)
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertEqual(paragraph.text, "Bukti: ")


if __name__ == "__main__":
    unittest.main()
