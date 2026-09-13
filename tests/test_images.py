"""Tests for shared JPEG/PNG/HEIC normalization."""

import io
import re
import unittest
from unittest.mock import patch

import requests
from docx import Document
from PIL import Image

from utils.images import (
    HAS_HEIF,
    HAS_PDF_RENDERER,
    MAX_WEB_IMAGE_BYTES,
    download_url_evidence,
    download_url_image,
    image_bytes_to_png,
    pdf_bytes_to_png_pages,
)


class _Response:
    """Small streamed response double used by the URL downloader tests."""

    def __init__(self, body=b"", *, status=200, content_type="image/jpeg",
                 content_length=None):
        self.content = body
        self.status_code = status
        self.headers = {"Content-Type": content_type}
        if content_length is not None:
            self.headers["Content-Length"] = str(content_length)
        self.cookies = {}
        self.closed = False

    def iter_content(self, chunk_size):
        for start in range(0, len(self.content), max(chunk_size // 2, 1)):
            yield self.content[start:start + max(chunk_size // 2, 1)]

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError(response=self)

    def close(self):
        self.closed = True


def _jpeg_bytes(size=(12, 8), *, orientation=None):
    source = io.BytesIO()
    image = Image.new("RGB", size, "blue")
    if orientation is None:
        image.save(source, format="JPEG")
    else:
        exif = Image.Exif()
        exif[274] = orientation
        image.save(source, format="JPEG", exif=exif)
    image.close()
    return source.getvalue()


def _png_bytes(size=(12, 8)):
    source = io.BytesIO()
    Image.new("RGB", size, "green").save(source, format="PNG")
    return source.getvalue()


class ImageNormalizationTests(unittest.TestCase):
    def test_jpeg_is_normalized_to_insertable_png(self):
        source = io.BytesIO()
        Image.new("RGB", (12, 8), "blue").save(source, format="JPEG")

        png_file, image = image_bytes_to_png(source.getvalue(), "image/jpeg")

        self.assertEqual(png_file.read(8), b"\x89PNG\r\n\x1a\n")
        self.assertEqual(image.size, (12, 8))
        png_file.seek(0)
        Document().add_paragraph().add_run().add_picture(png_file)

    @unittest.skipUnless(HAS_HEIF, "pillow-heif/native HEIF codec unavailable")
    def test_heic_is_decoded_to_insertable_png(self):
        source = io.BytesIO()
        Image.new("RGB", (11, 7), "green").save(source, format="HEIF")

        png_file, image = image_bytes_to_png(source.getvalue(), "image/heic")

        self.assertEqual(png_file.read(8), b"\x89PNG\r\n\x1a\n")
        self.assertEqual(image.size, (11, 7))
        png_file.seek(0)
        Document().add_paragraph().add_run().add_picture(png_file)

    def test_html_download_is_rejected_with_actionable_message(self):
        with self.assertRaisesRegex(RuntimeError, "halaman HTML"):
            image_bytes_to_png(b"<!doctype html><html></html>", "text/html")

    @unittest.skipUnless(HAS_PDF_RENDERER, "PyMuPDF unavailable")
    def test_multipage_pdf_is_rendered_to_insertable_png_pages(self):
        import pymupdf as fitz

        pdf = fitz.open()
        for text in ("Halaman satu", "Halaman dua"):
            page = pdf.new_page(width=300, height=500)
            page.insert_text((40, 60), text)
        raw_pdf = pdf.tobytes()
        pdf.close()

        pages = pdf_bytes_to_png_pages(raw_pdf)

        self.assertEqual(len(pages), 2)
        for stream, size in pages:
            self.assertEqual(stream.read(8), b"\x89PNG\r\n\x1a\n")
            self.assertGreater(size[0], 0)
            self.assertGreater(size[1], 0)
            stream.seek(0)
            Document().add_paragraph().add_run().add_picture(stream)


class RemoteImageSourceTests(unittest.TestCase):
    URL = "https://progres-mitra.bpskaro.site/api/dokumen-mitra/62/ss-fasih"

    def test_extensionless_jpeg_is_downloaded_and_normalized(self):
        response = _Response(_jpeg_bytes(), content_type="image/jpeg")
        with patch("utils.images.requests.get", return_value=response) as get:
            stream, image = download_url_image(self.URL)

        self.assertEqual(image.size, (12, 8))
        self.assertEqual(stream.read(8), b"\x89PNG\r\n\x1a\n")
        get.assert_called_once()
        self.assertTrue(response.closed)

    def test_extension_and_png_sources_are_supported(self):
        for url, body, content_type in (
            (
                "https://progres-mitra.bpskaro.site/api/dokumen-mitra/65/ss-fasih.jpeg",
                _jpeg_bytes(),
                "image/jpeg",
            ),
            ("https://example.test/api/image/123?download=1", _png_bytes(), "image/png"),
        ):
            with self.subTest(url=url):
                response = _Response(body, content_type=content_type)
                with patch("utils.images.requests.get", return_value=response):
                    stream, image = download_url_image(url)
                self.assertEqual(stream.read(8), b"\x89PNG\r\n\x1a\n")
                self.assertGreater(image.width, 0)
                image.close()

    def test_valid_image_wins_over_misleading_content_type(self):
        for content_type in (
            "application/octet-stream",
            "text/plain",
            "text/html",
            "application/pdf",
        ):
            with self.subTest(content_type=content_type):
                response = _Response(_jpeg_bytes(), content_type=content_type)
                with patch("utils.images.requests.get", return_value=response):
                    stream, image = download_url_evidence(self.URL)[0][1:]
                self.assertEqual(stream.read(8), b"\x89PNG\r\n\x1a\n")
                self.assertEqual(image, (12, 8))

    def test_request_streams_follows_redirects_and_keeps_tls_verification(self):
        response = _Response(_jpeg_bytes())
        with patch("utils.images.requests.get", return_value=response) as get:
            download_url_image("https://example.test/image")

        kwargs = get.call_args.kwargs
        self.assertTrue(kwargs["stream"])
        self.assertTrue(kwargs["allow_redirects"])
        self.assertTrue(kwargs["verify"])
        self.assertIn("timeout", kwargs)

    def test_html_response_is_rejected_even_with_binary_content_type(self):
        response = _Response(
            b"<!doctype html><html><body>login</body></html>",
            content_type="application/octet-stream",
        )
        with patch("utils.images.requests.get", return_value=response):
            with self.assertRaisesRegex(RuntimeError, "halaman HTML"):
                download_url_image(self.URL)
        self.assertTrue(response.closed)

    def test_empty_response_is_rejected(self):
        response = _Response(b"", content_type="image/jpeg")
        with patch("utils.images.requests.get", return_value=response):
            with self.assertRaisesRegex(RuntimeError, "respons kosong"):
                download_url_image(self.URL)

    def test_http_403_and_404_are_actionable_and_not_retried(self):
        for status, message in ((403, "Akses ditolak (403)"), (404, "File tidak ditemukan (404)")):
            with self.subTest(status=status):
                response = _Response(b"error", status=status)
                with patch("utils.images.requests.get", return_value=response) as get:
                    with self.assertRaisesRegex(RuntimeError, re.escape(message)):
                        download_url_image(self.URL)
                self.assertEqual(get.call_count, 1)

    def test_http_500_retries_then_succeeds(self):
        failed = _Response(b"temporary", status=500)
        success = _Response(_jpeg_bytes())
        with patch("utils.images.requests.get", side_effect=[failed, success]) as get, patch("utils.images.time.sleep") as sleep:
            stream, image = download_url_image(self.URL)

        self.assertEqual(stream.read(8), b"\x89PNG\r\n\x1a\n")
        image.close()
        self.assertEqual(get.call_count, 2)
        sleep.assert_called_once()
        self.assertTrue(failed.closed)
        self.assertTrue(success.closed)

    def test_timeout_retries_then_succeeds(self):
        success = _Response(_jpeg_bytes())
        with patch(
            "utils.images.requests.get",
            side_effect=[requests.Timeout("slow"), success],
        ) as get, patch("utils.images.time.sleep"):
            _, image = download_url_image(self.URL)

        image.close()
        self.assertEqual(get.call_count, 2)

    def test_download_size_limit_checks_header_before_reading_body(self):
        response = _Response(
            b"small",
            content_length=MAX_WEB_IMAGE_BYTES + 1,
        )
        with patch("utils.images.requests.get", return_value=response):
            with self.assertRaisesRegex(RuntimeError, "25 MB"):
                download_url_image(self.URL)
        self.assertTrue(response.closed)

    def test_corrupted_image_is_rejected_after_download(self):
        response = _Response(b"not an image", content_type="application/octet-stream")
        with patch("utils.images.requests.get", return_value=response):
            with self.assertRaisesRegex(RuntimeError, "bukan gambar yang didukung"):
                download_url_image(self.URL)

    def test_invalid_and_non_http_schemes_are_rejected(self):
        for url in ("", "file:///tmp/image.jpg", "ftp://example.test/image.jpg", "data:image/png;base64,AA="):
            with self.subTest(url=url):
                with self.assertRaisesRegex(ValueError, r"HTTP\(S\)"):
                    download_url_image(url)

    def test_google_drive_urls_keep_using_existing_downloader(self):
        sentinel = (io.BytesIO(b"png"), Image.new("RGB", (1, 1)))
        with patch("utils.images.download_drive_image", return_value=sentinel) as drive:
            result = download_url_image("https://drive.google.com/file/d/FILE_ID/view")

        drive.assert_called_once_with(
            "FILE_ID",
            timeout=15,
            max_retries=3,
            retry_delay=0.5,
            max_bytes=MAX_WEB_IMAGE_BYTES,
        )
        self.assertIs(result, sentinel)
        sentinel[1].close()

    def test_exif_orientation_is_applied_for_remote_image(self):
        response = _Response(_jpeg_bytes((2, 3), orientation=6))
        with patch("utils.images.requests.get", return_value=response):
            stream, image = download_url_image(self.URL)

        self.assertEqual(image.size, (3, 2))
        image.close()
        stream.close()

    @unittest.skipUnless(HAS_PDF_RENDERER, "PyMuPDF unavailable")
    def test_generic_pdf_is_rendered_for_evidence(self):
        import pymupdf as fitz

        pdf = fitz.open()
        pdf.new_page().insert_text((40, 60), "Bukti")
        raw_pdf = pdf.tobytes()
        pdf.close()
        response = _Response(raw_pdf, content_type="application/pdf")

        with patch("utils.images.requests.get", return_value=response):
            items = download_url_evidence("https://example.test/evidence")

        self.assertEqual(len(items), 1)
        self.assertEqual(items[0][0], "pdf_page")
        self.assertGreater(items[0][2][0], 0)


if __name__ == "__main__":
    unittest.main()
